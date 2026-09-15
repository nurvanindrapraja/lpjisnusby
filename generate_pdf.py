import os
import json
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable, Image
)
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#006837"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(36, 810, "LAPORAN PERTANGGUNGJAWABAN (LPJ) PC ISNU KOTA SURABAYA (2022 - 2026)")
            self.setStrokeColor(colors.HexColor("#bbf7d0"))
            self.setLineWidth(0.5)
            self.line(36, 802, 559, 802)
        
        # Footer (all pages)
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawString(36, 25, "PC ISNU Kota Surabaya | Website LPJ Resmi: https://isnusurabaya.or.id")
        page_str = f"Halaman {self._pageNumber} dari {page_count}"
        self.drawRightString(559, 25, page_str)
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(36, 35, 559, 35)
        
        self.restoreState()

def generate_pdf_charts():
    os.makedirs("Source", exist_ok=True)
    chart_paths = {}
    
    # 1. Seksi Bar Chart (Horizontal)
    fig, ax = plt.subplots(figsize=(5.5, 3.2), dpi=200)
    seksi_labels = [
        'Ideologi & Dakwah',
        'Pendidikan & Hum.',
        'Sains & Tekn. (Saintek)',
        'Kesehatan & Lingk.',
        'Perekonomian & Umat',
        'Organisasi & Kerjasama'
    ]
    seksi_values = [8, 9, 9, 8, 9, 10]
    bar_colors = ['#e11d48', '#d97706', '#7c3aed', '#0284c7', '#059669', '#006837']
    
    bars = ax.barh(seksi_labels, seksi_values, color=bar_colors, height=0.6, edgecolor='none')
    ax.set_title('Kegiatan per Seksi Bidang Kerja', fontsize=10, fontweight='bold', color='#004d28', pad=8)
    ax.set_xlim(0, 12)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['bottom'].set_color('#cbd5e1')
    ax.spines['left'].set_color('#cbd5e1')
    ax.xaxis.grid(True, linestyle='--', alpha=0.5, color='#e2e8f0')
    ax.set_axisbelow(True)
    ax.tick_params(axis='both', labelsize=8)
    
    for bar in bars:
        width = bar.get_width()
        ax.text(width + 0.3, bar.get_y() + bar.get_height()/2, f'{int(width)}',
                va='center', ha='left', fontsize=8, fontweight='bold', color='#0f172a')
                
    plt.tight_layout()
    chart_seksi_path = 'Source/chart_seksi.png'
    plt.savefig(chart_seksi_path, bbox_inches='tight', facecolor='white')
    plt.close()
    chart_paths['seksi'] = chart_seksi_path

    # 2. Peran Doughnut Chart
    fig, ax = plt.subplots(figsize=(4.5, 3.2), dpi=200)
    peran_labels = ['Penyelenggara\n(38)', 'Peserta\n(10)', 'Pembicara\n(5)']
    peran_sizes = [38, 10, 5]
    peran_colors = ['#006837', '#0284c7', '#7c3aed']
    
    wedges, texts, autotexts = ax.pie(
        peran_sizes, labels=peran_labels, colors=peran_colors, autopct='%1.1f%%',
        startangle=140, pctdistance=0.68,
        wedgeprops=dict(width=0.42, edgecolor='white', linewidth=2),
        textprops=dict(fontsize=7.5, color='#1e293b')
    )
    for autotext in autotexts:
        autotext.set_color('white')
        autotext.set_weight('bold')
        autotext.set_fontsize(7.5)
        
    ax.set_title('Peran ISNU dalam Kegiatan', fontsize=10, fontweight='bold', color='#004d28', pad=8)
    plt.tight_layout()
    chart_peran_path = 'Source/chart_peran.png'
    plt.savefig(chart_peran_path, bbox_inches='tight', facecolor='white')
    plt.close()
    chart_paths['peran'] = chart_peran_path

    # 3. Timeline Area Chart (Tren 2022 - 2026)
    fig, ax = plt.subplots(figsize=(8, 2.5), dpi=200)
    years = ['2022', '2023', '2024', '2025', '2026']
    counts = [3, 17, 15, 8, 10]
    
    ax.plot(years, counts, color='#006837', marker='o', linewidth=2.5, markersize=6, markerfacecolor='#d97706', markeredgecolor='white', markeredgewidth=1.5)
    ax.fill_between(years, counts, color='#006837', alpha=0.15)
    ax.set_title('Tren Grafik Jumlah Kegiatan Khidmah per Tahun (2022 - 2026)', fontsize=10, fontweight='bold', color='#004d28', pad=8)
    ax.set_ylim(0, 20)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['bottom'].set_color('#cbd5e1')
    ax.spines['left'].set_color('#cbd5e1')
    ax.yaxis.grid(True, linestyle='--', alpha=0.5, color='#e2e8f0')
    ax.set_axisbelow(True)
    ax.tick_params(axis='both', labelsize=8)
    
    for x, y in zip(years, counts):
        ax.annotate(f'{y} Kegiatan', (x, y), textcoords="offset points", xytext=(0, 6), ha='center', fontsize=8, fontweight='bold', color='#006837')

    plt.tight_layout()
    chart_timeline_path = 'Source/chart_timeline.png'
    plt.savefig(chart_timeline_path, bbox_inches='tight', facecolor='white')
    plt.close()
    chart_paths['timeline'] = chart_timeline_path

    return chart_paths

def build_pdf():
    pdf_filename = "Source/Dokumen_LPJ_PC_ISNU_Kota_Surabaya_2022-2026.pdf"
    os.makedirs("Source", exist_ok=True)
    
    # Generate charts
    chart_paths = generate_pdf_charts()

    doc = SimpleDocTemplate(
        pdf_filename,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=48,
        bottomMargin=48
    )

    styles = getSampleStyleSheet()
    
    # Custom Palette
    c_primary = colors.HexColor("#006837")
    c_dark = colors.HexColor("#004d28")
    c_gold = colors.HexColor("#b45309")
    c_text = colors.HexColor("#1e293b")
    c_bg_light = colors.HexColor("#f0fdf4")
    
    # Custom Paragraph Styles
    style_cover_title = ParagraphStyle(
        'CoverTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=c_primary,
        alignment=1, # Center
        spaceAfter=8
    )
    
    style_cover_subtitle = ParagraphStyle(
        'CoverSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=c_dark,
        alignment=1,
        spaceAfter=12
    )

    style_cover_meta = ParagraphStyle(
        'CoverMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
        alignment=1,
        spaceAfter=15
    )

    style_h1 = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=c_dark,
        spaceBefore=12,
        spaceAfter=8,
        keepWithNext=True
    )

    style_h2 = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=14,
        textColor=c_primary,
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True
    )

    style_body = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13.5,
        textColor=c_text,
        spaceAfter=6
    )

    style_table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=c_text
    )

    style_table_cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        textColor=c_text
    )

    style_table_cell_header = ParagraphStyle(
        'TableCellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        textColor=colors.white,
        alignment=1
    )

    story = []

    # ------------------ COVER / HEADER TITLE ------------------
    story.append(Paragraph("PIMPINAN CABANG IKATAN SARJANA NAHDLATUL ULAMA KOTA SURABAYA", ParagraphStyle('TopHeader', fontName='Helvetica-Bold', fontSize=10, leading=12, textColor=c_primary, alignment=1)))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceAfter=12))
    
    story.append(Paragraph("RANGKUMAN LAPORAN PERTANGGUNGJAWABAN (LPJ)", style_cover_title))
    story.append(Paragraph("Masa Khidmat 2022 – 2026", style_cover_subtitle))
    story.append(Paragraph("<b>Kesekretariatan:</b> Jl. Bubutan Gg. VI No.2, Alun-alun Contong, Bubutan, Kota Surabaya<br/><b>Kontak:</b> 081330278721 | <b>Email:</b> isnukotasurabaya@gmail.com | <b>Website:</b> https://isnusurabaya.or.id", style_cover_meta))
    
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#bbf7d0"), spaceAfter=12))

    # ------------------ BAB I: RINGKASAN EKSEKUTIF & VISUALISASI GRAFIK ------------------
    story.append(Paragraph("BAB I: RINGKASAN EKSEKUTIF & VISUALISASI GRAFIK KHIDMAH", style_h1))
    story.append(Paragraph(
        "Pimpinan Cabang Ikatan Sarjana Nahdlatul Ulama (PC ISNU) Kota Surabaya Masa Khidmat 2022–2026 telah melaksanakan seluruh agenda jam'iyyah dan pengabdian masyarakat secara berkelanjutan. Laporan ini merangkum seluruh rekam jejak kegiatan, regulasi organisasi, publikasi riset ilmiah, prestasi PWNU Jatim Award, serta transparansi pengelolaan kas keuangan.",
        style_body
    ))
    
    # Stat Metric Cards
    stat_data = [
        [
            Paragraph("<b>53</b><br/><font size=7 color='#475569'>TOTAL KEGIATAN</font>", ParagraphStyle('C1', alignment=1, fontSize=11, leading=13, fontName='Helvetica-Bold', textColor=c_primary)),
            Paragraph("<b>6</b><br/><font size=7 color='#475569'>SEKSI BIDANG KERJA</font>", ParagraphStyle('C2', alignment=1, fontSize=11, leading=13, fontName='Helvetica-Bold', textColor=c_gold)),
            Paragraph("<b>38</b><br/><font size=7 color='#475569'>PENYELENGGARA</font>", ParagraphStyle('C3', alignment=1, fontSize=11, leading=13, fontName='Helvetica-Bold', textColor=colors.HexColor("#1d4ed8"))),
            Paragraph("<b>5 / 10</b><br/><font size=7 color='#475569'>PEMBICARA / PESERTA</font>", ParagraphStyle('C4', alignment=1, fontSize=11, leading=13, fontName='Helvetica-Bold', textColor=colors.HexColor("#7e22ce")))
        ]
    ]
    t_stat = Table(stat_data, colWidths=[125, 125, 125, 148])
    t_stat.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_bg_light),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#bbf7d0")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t_stat)
    story.append(Spacer(1, 10))

    # Visual Charts Section (Side-by-Side: Seksi Bar Chart & Peran Doughnut Chart)
    story.append(Paragraph("Visualisasi Grafik Capaian Kegiatan & Peran ISNU", style_h2))
    
    img_seksi = Image(chart_paths['seksi'], width=265, height=155)
    img_peran = Image(chart_paths['peran'], width=245, height=155)
    
    chart_grid = Table([[img_seksi, img_peran]], colWidths=[270, 253])
    chart_grid.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('PADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(chart_grid)
    story.append(Spacer(1, 8))

    # Full Width Timeline Chart
    img_timeline = Image(chart_paths['timeline'], width=523, height=160)
    story.append(img_timeline)
    story.append(Spacer(1, 12))

    # Seksi Distribution Summary Table
    seksi_summary = [
        [Paragraph("Seksi Bidang Kerja", style_table_cell_header), Paragraph("Jumlah Kegiatan Khidmah", style_table_cell_header)],
        [Paragraph("1. Seksi Organisasi & Kerjasama Antar Lembaga", style_table_cell), Paragraph("10 Kegiatan", style_table_cell_bold)],
        [Paragraph("2. Seksi Perekonomian & Pemberdayaan Umat", style_table_cell), Paragraph("9 Kegiatan", style_table_cell_bold)],
        [Paragraph("3. Seksi Kesehatan Masyarakat & Lingkungan Hidup", style_table_cell), Paragraph("8 Kegiatan", style_table_cell_bold)],
        [Paragraph("4. Seksi Sains & Teknologi (SAINTEK)", style_table_cell), Paragraph("9 Kegiatan", style_table_cell_bold)],
        [Paragraph("5. Seksi Pendidikan & Humaniora", style_table_cell), Paragraph("9 Kegiatan", style_table_cell_bold)],
        [Paragraph("6. Seksi Penguatan Ideologi & Dakwah Digital", style_table_cell), Paragraph("8 Kegiatan", style_table_cell_bold)],
        [Paragraph("<b>TOTAL REKAPITULASI KEGIATAN</b>", style_table_cell), Paragraph("<b>53 KEGIATAN</b>", style_table_cell_bold)]
    ]
    t_seksi = Table(seksi_summary, colWidths=[380, 143])
    t_seksi.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_dark),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#e6f4ed")),
        ('PADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_seksi)
    story.append(Spacer(1, 15))

    # ------------------ BAB II: STRUKTUR ORGANISASI & REGULASI ------------------
    story.append(Paragraph("BAB II: STRUKTUR ORGANISASI & DOKUMEN REGULASI", style_h1))
    
    story.append(Paragraph("1. Surat Keputusan (SK) Legalitas Organisasi", style_h2))
    sk_table_data = [
        [Paragraph("No / Kategori", style_table_cell_header), Paragraph("Judul & Nomor Surat Keputusan", style_table_cell_header), Paragraph("Keterangan", style_table_cell_header)],
        [Paragraph("1. Pimpinan Cabang", style_table_cell_bold), Paragraph("SK PC ISNU Kota Surabaya 2022-2026<br/><font color='#64748b'>No: 235/SK/PP-ISNU/XI/2022</font>", style_table_cell), Paragraph("SK Penetapan Susunan Pengurus PC ISNU Kota Surabaya oleh PP ISNU", style_table_cell)],
        [Paragraph("2. Tim Kerja", style_table_cell_bold), Paragraph("SK Tim Pengelola Website & Medsos<br/><font color='#64748b'>No: 13/SK/PC-ISNU/XII/2022</font>", style_table_cell), Paragraph("SK Pembentukan Tim IT & Pengelola Media Informasi", style_table_cell)],
        [Paragraph("3. Panitia", style_table_cell_bold), Paragraph("SK Panitia Konfercab VI ISNU Surabaya<br/><font color='#64748b'>No: 9/SK/PC-ISNU/VII/2026</font>", style_table_cell), Paragraph("SK Penetapan Panitia Konferensi Cabang VI", style_table_cell)],
        [Paragraph("4. PAC (Kecamatan)", style_table_cell_bold), Paragraph("SK PAC ISNU Gunung Anyar, Gayungan, Lakarsantri, Sukolilo, Semampir, Wonokromo", style_table_cell), Paragraph("SK Pengesahan Pengurus Anak Cabang (PAC) Se-Kota Surabaya", style_table_cell)],
    ]
    t_sk = Table(sk_table_data, colWidths=[110, 240, 173])
    t_sk.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('PADDING', (0,0), (-1,-1), 4.5),
    ]))
    story.append(t_sk)
    story.append(Spacer(1, 10))

    story.append(Paragraph("2. Roadmap & Renstra ISNU Surabaya 2022–2026 Menuju Indonesia Emas 2045", style_h2))
    story.append(Paragraph("Roadmap ini meletakkan pondasi bagi transformasi sarjana Nahdliyin Kota Surabaya melalui <b>4 Pilar Utama Strategi Khidmah</b>:", style_body))
    pillars = [
        "<b>• Penguatan Kapasitas SDM Akademik & Riset:</b> Publikasi jurnal bereputasi, studi lanjut, dan kualifikasi Guru Besar kader NU.",
        "<b>• Transformasi Digital & Dakwah Siber:</b> Penguasaan sains data, Generative AI, dan penguatan narasi moderasi beragama An-Nahdliyah.",
        "<b>• Kemandirian Ekonomi & Industri Halal:</b> Pendampingan sertifikasi halal UMKM, penciptaan Juleha kompeten, dan QRIS digital.",
        "<b>• Sinergi Kelembagaan & Pengabdian Masyarakat:</b> Kemitraan strategis Perguruan Tinggi (ITS, UNUSA, UNESA, UNISMA) dan khitanan massal."
    ]
    for p in pillars:
        story.append(Paragraph(p, style_body))
    story.append(Spacer(1, 10))

    story.append(Paragraph("3. Standard Operating Procedure (SOP) Ber-TTD Resmi", style_h2))
    story.append(Paragraph("PC ISNU Kota Surabaya menetapkan 3 dokumen Standar Operasional Prosedur (SOP) resmi yaitu: <b>(1) SOP Kesekretariatan & Persuratan</b>, <b>(2) SOP Kebendaharaan & Keuangan (Laporan 3 Bulanan & Pencairan Proposal)</b>, serta <b>(3) SOP Inventarisasi Aset Organisasi</b>.", style_body))
    
    story.append(Spacer(1, 15))

    # ------------------ BAB III: PWNU JATIM AWARD 2023 ------------------
    story.append(Paragraph("BAB III: CAPAIAN PRESTASI PWNU JATIM AWARD 2023", style_h1))
    story.append(Paragraph(
        "<b>PC ISNU Surabaya Raih Juara 2 Kategori Banom NU Surabaya Dalam PWNU Jatim Award 2023</b><br/>"
        "Pengurus Cabang Ikatan Sarjana Nahdlatul Ulama (PC ISNU) Kota Surabaya berhasil meraih <b>Juara 2</b> dalam kategori Badan Otonom (Banom) NU Surabaya pada ajang penganugerahan PWNU Jatim Award 2023 di Kediri. Penghargaan ini diraih atas dasar keberanian melahirkan berbagai inovasi program khidmah yang kolaboratif tingkat lokal, nasional, dan internasional.",
        style_body
    ))
    quote_data = [[
        Paragraph("<i>\"Sangat bersyukur atas penghargaan dalam ajang PWNU Award. Capaian ini bukan hasil kerja individu, namun kerja bersama semua pengurus dalam beberapa tahun terakhir. Semoga menjadi awal untuk terus berkhidmat di ISNU dalam memberikan manfaat lebih besar untuk warga Surabaya, khususnya warga nahdliyyin.\"</i><br/><br/><b>- Ahmad Bashri (Ketua PC ISNU Kota Surabaya)</b>", ParagraphStyle('QStyle', parent=style_body, textColor=c_dark, leading=14))
    ]]
    t_quote = Table(quote_data, colWidths=[523])
    t_quote.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), c_bg_light),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#bbf7d0")),
        ('PADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_quote)
    story.append(Spacer(1, 15))

    # ------------------ BAB IV: PUBLIKASI ILMIAH JOSHNU ------------------
    story.append(Paragraph("BAB IV: PUBLIKASI ILMIAH JOSHNU", style_h1))
    story.append(Paragraph(
        "<b>Journal of Sustainable Halal Industry and New Urbanism (JOSHNU)</b> merupakan jurnal akademis terdepan PC ISNU Kota Surabaya (Rilis Perdana Vol. 1 2025) yang mengeksplorasi industri halal berkelanjutan dan prinsip urbanisme baru era Society 5.0. 6 Rumpun Fokus Utama JOSHNU meliputi: Makanan & Minuman Halal, Modest & Fashion, Pariwisata Halal, Farmasi & Kosmetika Halal, Media & Rekreasi Syariah, serta Bisnis Syariah & Urbanism.",
        style_body
    ))
    story.append(Spacer(1, 15))

    # ------------------ BAB V: REKAPITULASI 53 KEGIATAN ------------------
    story.append(PageBreak()) # Clean page break for table of activities
    story.append(Paragraph("BAB V: REKAPITULASI RINCIAN 53 KEGIATAN (2022 - 2026)", style_h1))
    story.append(Paragraph("Berikut adalah daftar rekapitulasi lengkap 53 kegiatan khidmah PC ISNU Kota Surabaya beserta Output dan Outcome:", style_body))

    # Load activities from excel
    wb = openpyxl.load_workbook('Daftar Kegiatan ISNU Kota Surabaya 2022 - 2026.xlsx', data_only=True)
    sheet = wb.active
    
    act_table_data = [
        [
            Paragraph("No", style_table_cell_header),
            Paragraph("Tanggal & Nama Kegiatan", style_table_cell_header),
            Paragraph("Peran & Seksi", style_table_cell_header),
            Paragraph("Output & Outcome Kegiatan", style_table_cell_header)
        ]
    ]

    for r in range(2, 55):
        no = sheet.cell(row=r, column=1).value
        tgl_val = sheet.cell(row=r, column=2).value
        tgl_str = str(tgl_val)[:10] if tgl_val else '-'
        kegiatan = sheet.cell(row=r, column=3).value or ''
        lokasi = sheet.cell(row=r, column=4).value or 'Surabaya'
        peran = sheet.cell(row=r, column=5).value or 'Penyelenggara'
        seksi1 = sheet.cell(row=r, column=6).value or ''
        output = sheet.cell(row=r, column=8).value or '-'
        outcome = sheet.cell(row=r, column=9).value or '-'

        c_info = f"<b>{kegiatan}</b><br/><font color='#64748b'>📅 {tgl_str} | 📍 {lokasi}</font>"
        c_peran = f"<b>{peran}</b><br/><font color='#006837'>{seksi1}</font>"
        c_out = f"<b>Output:</b> {output}<br/><b>Outcome:</b> {outcome}"

        act_table_data.append([
            Paragraph(str(no), ParagraphStyle('CNo', alignment=1, fontSize=8, fontName='Helvetica-Bold')),
            Paragraph(c_info, style_table_cell),
            Paragraph(c_peran, style_table_cell),
            Paragraph(c_out, style_table_cell)
        ])

    t_act = Table(act_table_data, colWidths=[25, 170, 115, 213])
    t_act.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('PADDING', (0,0), (-1,-1), 4),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))
    story.append(t_act)
    story.append(Spacer(1, 15))

    # ------------------ BAB VI: LAPORAN KEUANGAN ------------------
    story.append(KeepTogether([
        Paragraph("BAB VI: LAPORAN KEUANGAN & AKUNTABILITAS KAS", style_h1),
        Paragraph("Rekapitulasi keuangan PC ISNU Kota Surabaya Periode 2022–2026 menyajikan pengelolaan kas masuk, pengeluaran program, serta saldo kas akhir secara rinci dan teraudit:", style_body)
    ]))

    fin_stat_data = [
        [
            Paragraph("<b>Rp 26.976.000</b><br/><font size=7 color='#15803d'>TOTAL PEMASUKAN KAS</font>", ParagraphStyle('F1', alignment=1, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#15803d"))),
            Paragraph("<b>Rp 20.676.000</b><br/><font size=7 color='#dc2626'>TOTAL PENGELUARAN PROGRAM</font>", ParagraphStyle('F2', alignment=1, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#dc2626"))),
            Paragraph("<b>Rp 6.300.000</b><br/><font size=7 color='#1d4ed8'>SALDO KAS AKHIR</font>", ParagraphStyle('F3', alignment=1, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#1d4ed8"))),
            Paragraph("<b>76,6%</b><br/><font size=7 color='#7e22ce'>EFISIENSI KHIDMAH</font>", ParagraphStyle('F4', alignment=1, fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor("#7e22ce")))
        ]
    ]
    t_fin_stat = Table(fin_stat_data, colWidths=[130, 130, 130, 133])
    t_fin_stat.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f8fafc")),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(t_fin_stat)
    story.append(Spacer(1, 10))

    fin_trans_data = [
        [
            Paragraph("No / Tgl", style_table_cell_header),
            Paragraph("Uraian Transaksi Keuangan", style_table_cell_header),
            Paragraph("Pemasukan (Debet)", style_table_cell_header),
            Paragraph("Pengeluaran (Kredit)", style_table_cell_header)
        ],
        [Paragraph("1. 25/01/2023", style_table_cell), Paragraph("Sisa Konfercab V", style_table_cell), Paragraph("Rp 2.324.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("2. 10/01/2023", style_table_cell), Paragraph("Pemasukan Iuran Sukarela Pengurus", style_table_cell), Paragraph("Rp 1.000.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("3. 10/01/2023", style_table_cell), Paragraph("Pemasukan Iuran Wajib Pengurus", style_table_cell), Paragraph("Rp 4.100.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("4. 12/03/2023", style_table_cell), Paragraph("Karangan Bunga untuk Bpk. Prof. Jadid", style_table_cell), Paragraph("-", style_table_cell), Paragraph("Rp 500.000", style_table_cell_bold)],
        [Paragraph("5. 17/05/2023", style_table_cell), Paragraph("Tarik Tunai Ketua ISNU untuk Operasional Kegiatan", style_table_cell), Paragraph("-", style_table_cell), Paragraph("Rp 3.000.000", style_table_cell_bold)],
        [Paragraph("6. 16/05/2025", style_table_cell), Paragraph("Donatur dari Bapak Mahirul Mursid", style_table_cell), Paragraph("Rp 5.000.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("7. 01/01/2025", style_table_cell), Paragraph("Pemasukan Iuran Wajib Pengurus", style_table_cell), Paragraph("Rp 5.700.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("8. 01/01/2025", style_table_cell), Paragraph("Donatur dari Bapak Mahirul Mursid", style_table_cell), Paragraph("Rp 5.000.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("9. 01/01/2025", style_table_cell), Paragraph("Donatur dari Rektor UPN", style_table_cell), Paragraph("Rp 1.000.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("10. -", style_table_cell), Paragraph("Saldo Pengisian Kas Bu Firdaus", style_table_cell), Paragraph("Rp 30.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("11. -", style_table_cell), Paragraph("Saldo Pengisian Kas Pak Ketua", style_table_cell), Paragraph("Rp 322.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("12. 08/01/2025", style_table_cell), Paragraph("Donasi dr. Abraham Ali Firdaus", style_table_cell), Paragraph("Rp 2.500.000", style_table_cell_bold), Paragraph("-", style_table_cell)],
        [Paragraph("13. 03/01/2025", style_table_cell), Paragraph("Pelaksanaan Program Sunatan Massal Gratis UNUSA", style_table_cell), Paragraph("-", style_table_cell), Paragraph("Rp 17.176.000", style_table_cell_bold)],
        [Paragraph("<b>TOTAL REKAPITULASI</b>", style_table_cell_bold), Paragraph("<b>AKUMULASI KAS KHIDMAH</b>", style_table_cell_bold), Paragraph("<b>Rp 26.976.000</b>", ParagraphStyle('TB1', parent=style_table_cell_bold, textColor=colors.HexColor("#15803d"))), Paragraph("<b>Rp 20.676.000</b>", ParagraphStyle('TB2', parent=style_table_cell_bold, textColor=colors.HexColor("#dc2626")))],
        [Paragraph("<b>SALDO KAS AKHIR</b>", style_table_cell_bold), Paragraph("<b>SALDO KAS PC ISNU KOTA SURABAYA</b>", style_table_cell_bold), Paragraph("<b>Rp 6.300.000</b>", ParagraphStyle('TB3', parent=style_table_cell_bold, textColor=colors.HexColor("#1d4ed8"))), Paragraph("<b>TERALOKASI 100%</b>", style_table_cell_bold)]
    ]
    t_fin_trans = Table(fin_trans_data, colWidths=[80, 233, 105, 105])
    t_fin_trans.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_dark),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('BACKGROUND', (0,-2), (-1,-2), colors.HexColor("#f1f5f9")),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor("#e6f4ed")),
        ('PADDING', (0,0), (-1,-1), 4.5),
    ]))
    story.append(t_fin_trans)

    # Build document
    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF generated with embedded charts successfully: {pdf_filename}")

if __name__ == "__main__":
    build_pdf()
